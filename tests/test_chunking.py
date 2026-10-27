"""Structure-aware chunk boundaries and source-offset regression tests."""
import pytest

import config
from rag_store import RagStore


@pytest.mark.parametrize("heading", ["# New section", "## New section", "New section\n-----------"])
@pytest.mark.parametrize("overlap", [0, 20])
def test_splits_before_heading_not_after_it(monkeypatch, heading, overlap):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 100)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", overlap)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    text = "Previous paragraph. " * 4 + "\n\n" + heading + "\n\nFollowing text block."
    chunks = RagStore.__new__(RagStore)._chunk_document(text)
    assert len(chunks) >= 2
    for content, start, end in chunks:
        assert content == text[start:end].strip()
        if "New section" in content:
            assert "Following text block." in content
    assert "New section" not in chunks[0][0]
    assert chunks[-1][2] == len(text)


def test_heading_and_first_block_stay_together(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 120)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 15)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    body = "The following paragraph belongs to the heading."
    text = "Previous text. " * 5 + "\n\n# Heading\n\n" + body + "\n\n" + "Later text. " * 15
    chunks = RagStore.__new__(RagStore)._chunk_document(text)
    assert any("# Heading\n\n" + body in content for content, _, _ in chunks)
    assert all(body in content for content, _, _ in chunks if "# Heading" in content)


def test_oversized_heading_block_makes_progress_without_losing_text(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 45)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 5)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    text = "# Heading\n\n" + "body " * 80 + "\n\n## Next\n\nLast paragraph."
    chunks = RagStore.__new__(RagStore)._chunk_document(text)
    assert len(chunks) < len(text)
    assert all(end - start <= 45 for _, start, end in chunks)
    assert chunks[-1][2] == len(text)
    assert all(next_start <= end for (_, _, end), (_, next_start, _) in zip(chunks, chunks[1:]))
    assert all(content == text[start:end].strip() for content, start, end in chunks)


def test_plain_text_and_empty_input(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 80)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 10)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    chunker = RagStore.__new__(RagStore)
    assert chunker._chunk_document(" \n") == []
    assert chunker._chunk_document("Short paragraph.") == [("Short paragraph.", 0, 16)]
    text = "A plain paragraph. " * 30
    chunks = chunker._chunk_document(text)
    assert chunks[0][1] == 0
    assert chunks[-1][2] == len(text)
    assert all(content == text[start:end].strip() for content, start, end in chunks)


def test_short_document_is_not_split_at_heading(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 150)
    text = "Introduction paragraph.\n\n# Heading\n\nA complete following paragraph."
    assert RagStore.__new__(RagStore)._chunk_document(text) == [(text, 0, len(text))]


def test_consecutive_headings_stay_with_following_block(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 100)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 15)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    block = "# Title\n\n## Subheading\n\nThe associated body."
    text = "Previous paragraph. " * 4 + "\n\n" + block
    chunks = RagStore.__new__(RagStore)._chunk_document(text)
    assert chunks[-1][0] == block
    assert all("#" not in content for content, _, _ in chunks[:-1])


def test_code_fence_hashes_are_not_headings(monkeypatch):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 100)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 0)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    prefix = "Previous paragraph. " * 4 + "\n\n```python\n"
    text = prefix + "# code comment\n\nprint('hello')\n```\n\n" + "Later text. " * 15
    chunks = RagStore.__new__(RagStore)._chunk_document(text)
    # No special split is introduced immediately before this comment.
    assert len(prefix) not in [end for _, _, end in chunks]


@pytest.mark.parametrize("chunk_size,overlap,margin", [(1, 0, 0), (40, 35, 0), (100, 20, 15), (180, 0, 30)])
def test_heading_boundaries_preserve_coverage_and_bounds(monkeypatch, chunk_size, overlap, margin):
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", chunk_size)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", overlap)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", margin)
    text = ("Intro. " * 13 + "\n\n# First\n\n" + "First body. " * 8
            + "\n\n## Second\n\nShort body.\n\n" + "Final text. " * 20)
    chunks = RagStore.__new__(RagStore)._chunk_document(text)
    covered = set()
    previous_start = -1
    for content, start, end in chunks:
        assert start > previous_start
        assert 0 < end - start <= chunk_size + margin
        assert content == text[start:end].strip()
        covered.update(range(start, end))
        previous_start = start
    assert all(index in covered for index, character in enumerate(text) if not character.isspace())


def test_docx_headings_reach_persisted_chunks(data_env, tmp_path, monkeypatch):
    from docx import Document
    from document_loader import load_document

    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 110)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 15)
    monkeypatch.setattr(config, "CHUNK_SPLIT_MARGIN", 0)
    document = Document()
    document.add_paragraph("Introduction paragraph. " * 5)
    document.add_heading("Contract conditions", level=1)
    body = "Cancellation requires thirty days of notice."
    document.add_paragraph(body)
    path = tmp_path / "contract.docx"
    document.save(path)
    loaded = load_document(path)
    store = RagStore("heading-integration")
    try:
        store.add_documents([loaded.content], [loaded.metadata])
        chunks = store.get_chunks(source=str(loaded.path), limit=100)
        assert any("# Contract conditions\n\n" + body in chunk["content"] for chunk in chunks)
        assert all(body in chunk["content"] for chunk in chunks if "# Contract conditions" in chunk["content"])
    finally:
        store.close()
