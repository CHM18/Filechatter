"""Tests for document_loader: new readers and the file-type catalog."""
from __future__ import annotations

from pathlib import Path

import pytest

import document_loader as dl


def write(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


class TestTextLikeReaders:
    @pytest.mark.parametrize("name", ["notes.md", "script.py", "data.csv", "config.yaml", "page.xml"])
    def test_reads_as_text(self, tmp_path, name):
        path = write(tmp_path, name, "hello world content")
        loaded = dl.load_document(path)
        assert "hello world content" in loaded.content

    def test_markdown_routes_to_documentation(self, tmp_path):
        path = write(tmp_path, "readme.md", "# Title")
        assert dl.load_document(path).metadata["collection_name"] == "documentation"

    def test_code_routes_to_code(self, tmp_path):
        path = write(tmp_path, "main.py", "print('hi')")
        assert dl.load_document(path).metadata["collection_name"] == "code"

    def test_html_strips_tags(self, tmp_path):
        path = write(tmp_path, "page.html",
                     "<html><head><style>x{}</style></head><body><h1>Hi</h1><p>There &amp; more</p></body></html>")
        content = dl.load_document(path).content
        assert "Hi" in content and "There & more" in content
        assert "<h1>" not in content
        assert "x{}" not in content  # style stripped


class TestOfficeReaders:
    def test_docx_preserves_blocks_textboxes_and_controls(self, tmp_path):
        from docx import Document
        from docx.oxml import parse_xml

        document = Document()
        document.add_heading("Contract", 0)
        document.add_paragraph("Introduction")
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Table before textbox"
        document.element.body.insert(-1, parse_xml("""
            <w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                   xmlns:v="urn:schemas-microsoft-com:vml"
                   xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
                   xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
                   xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">
              <w:sdtContent>
                <w:p><w:r><w:t>Control value</w:t></w:r></w:p>
                <w:p><w:r><w:t>Anchor</w:t>
                  <mc:AlternateContent>
                    <mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx>
                      <w:txbxContent>
                        <w:p><w:r><w:t>Modern textbox</w:t></w:r></w:p>
                        <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Textbox table</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
                      </w:txbxContent>
                    </wps:txbx></wps:wsp></w:drawing></mc:Choice>
                    <mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>
                      <w:p><w:r><w:t>Duplicate fallback</w:t></w:r></w:p>
                    </w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>
                  </mc:AlternateContent>
                </w:r></w:p>
                <w:p><w:r><w:pict><v:shape><v:textbox><w:txbxContent>
                  <w:p><w:r><w:t>Legacy textbox</w:t></w:r></w:p>
                </w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>
                <w:p><w:sdt><w:sdtPr><w14:checkbox><w14:checked w14:val="1"/></w14:checkbox></w:sdtPr>
                  <w:sdtContent><w:r><w:t>Approved</w:t></w:r></w:sdtContent>
                </w:sdt></w:p>
              </w:sdtContent>
            </w:sdt>
        """))
        document.add_paragraph("Conclusion")
        path = tmp_path / "controls.docx"
        document.save(path)

        text = dl.load_document(path).content
        expected = ["# Contract", "Introduction", "Table before textbox", "Control value",
                    "Anchor", "Modern textbox", "Textbox table", "Legacy textbox",
                    "[x] Approved", "Conclusion"]
        positions = [text.index(value) for value in expected]
        assert positions == sorted(positions)
        assert text.count("Modern textbox") == 1
        assert "Duplicate fallback" not in text

    def test_docx_inherited_headings_and_shared_headers(self, tmp_path):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE

        document = Document()
        style = document.styles.add_style("Custom section", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Heading 2"]
        document.add_paragraph("Inherited heading", style=style)
        document.add_paragraph("Section body")
        document.sections[0].header.paragraphs[0].text = "Shared header"
        document.sections[0].footer.paragraphs[0].text = "Shared footer"
        document.add_section()
        path = tmp_path / "sections.docx"
        document.save(path)

        text = dl.load_document(path).content
        assert "## Inherited heading\n\nSection body" in text
        assert text.count("Shared header") == 1
        assert text.count("Shared footer") == 1

    def test_docx_inline_content_fields_and_fallback(self, tmp_path):
        from docx import Document
        from docx.oxml import parse_xml

        document = Document()
        document.element.body.insert(-1, parse_xml("""
          <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
               xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
               xmlns:unsupported="urn:unsupported"
               xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">
            <w:r><w:t>Start</w:t><w:tab/></w:r>
            <w:hyperlink><w:r><w:t>Link text</w:t><w:br/></w:r></w:hyperlink>
            <w:r><w:instrText>PRIVATE FIELD INSTRUCTION</w:instrText></w:r>
            <w:r><w:t>Stored field result</w:t></w:r>
            <w:del><w:r><w:delText>Deleted text</w:delText></w:r></w:del>
            <w:ins><w:r><w:t> Inserted text</w:t></w:r></w:ins>
            <w:r><w:fldChar><w:ffData><w:checkBox><w:default w:val="1"/></w:checkBox></w:ffData></w:fldChar></w:r>
            <w:sdt><w:sdtPr><w14:checkbox><w14:checked w14:val="0"/></w14:checkbox></w:sdtPr>
              <w:sdtContent><w:r><w:t>&#x2610;</w:t></w:r></w:sdtContent>
            </w:sdt>
            <mc:AlternateContent>
              <mc:Choice Requires="unsupported"><w:r><w:t>Unsupported choice</w:t></w:r></mc:Choice>
              <mc:Fallback><w:r><w:t>Fallback text</w:t></w:r></mc:Fallback>
            </mc:AlternateContent>
          </w:p>
        """))
        path = tmp_path / "fields.docx"
        document.save(path)

        text = dl.load_document(path).content
        assert "Start\tLink text\nStored field result Inserted text[x] [ ] Fallback text" in text
        assert "PRIVATE FIELD" not in text
        assert "Deleted text" not in text
        assert "Unsupported choice" not in text

    def test_docx_ocr_uses_table_and_header_image_relationships(self, tmp_path, monkeypatch):
        import io
        from docx import Document
        from PIL import Image

        document = Document()
        buffer = io.BytesIO()
        Image.new("RGB", (10, 10), color="white").save(buffer, format="PNG")
        for paragraph in (document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0],
                          document.sections[0].header.paragraphs[0]):
            buffer.seek(0)
            paragraph.add_run().add_picture(buffer)
        path = tmp_path / "table-images.docx"
        document.save(path)
        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        calls = []

        def fake_ocr(data):
            calls.append(data)
            return "Extracted picture text"

        monkeypatch.setattr(dl, "_ocr_image_bytes", fake_ocr)
        loaded = dl.load_document(path, ocr_enabled=True)
        assert len(calls) == 2
        assert loaded.content.count("[Image OCR]: Extracted picture text") == 2
        assert loaded.metadata["ocr_image_count"] == 2
        calls.clear()
        assert "Image OCR" not in dl.load_document(path, ocr_enabled=False).content
        assert calls == []

    def test_reads_xlsx(self, tmp_path):
        import openpyxl

        path = tmp_path / "book.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Name", "Score"])
        ws.append(["Ada", 99])
        wb.save(str(path))

        content = dl.load_document(path).content
        assert "Sheet: Data" in content
        assert "Name | Score" in content
        assert "Ada" in content and "99" in content

    def test_reads_pptx(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches

        path = tmp_path / "deck.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "Quarterly results summary"
        prs.save(str(path))

        content = dl.load_document(path).content
        assert "Slide 1" in content
        assert "Quarterly results summary" in content


class TestSupportAndCatalog:
    def test_supported_extensions_expanded(self):
        for ext in [".md", ".py", ".html", ".csv", ".xlsx", ".pptx", ".json"]:
            assert ext in dl.SUPPORTED_EXTENSIONS

    def test_catalog_has_categories(self):
        catalog = dl.file_type_catalog()
        names = {c["category"] for c in catalog}
        assert {"Office", "Images", "Text", "Web", "Code"}.issubset(names)

    def test_catalog_reports_availability(self):
        catalog = {c["category"]: c for c in dl.file_type_catalog()}
        office = {e["ext"]: e for e in catalog["Office"]["extensions"]}
        assert office[".xlsx"]["available"] is True  # openpyxl installed in test env
        text = {e["ext"]: e for e in catalog["Text"]["extensions"]}
        assert text[".md"]["available"] is True
        assert text[".md"]["dependency"] is None

    def test_support_status_marks_text_ready(self):
        support = dl.get_document_support_status()
        assert support[".md"]["available"] is True
        assert support[".py"]["category"] == "Code"

    def test_image_uses_vision_and_exif_pipeline(self, tmp_path, monkeypatch):
        path = tmp_path / "image.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n")

        monkeypatch.setattr(
            dl,
            "_summarize_image_with_llm",
            lambda _path, **_kw: "A scenic mountain view.\nSearchable keywords: mountain, alpine lake, hiking",
        )
        monkeypatch.setattr(dl, "_extract_image_exif", lambda _path: {"Make": "Canon"})

        loaded = dl.load_document(path)
        assert "A scenic mountain view." in loaded.content
        assert "Searchable keywords: mountain, alpine lake, hiking" in loaded.content
        assert loaded.metadata["content_kind"] == "image"
        assert loaded.metadata["image_summary"] == "A scenic mountain view."
        assert loaded.metadata["image_keywords"] == ["mountain", "alpine lake", "hiking"]
        assert loaded.metadata["image_exif"]["Make"] == "Canon"

    def test_vision_corrects_exif_orientation_without_changing_source(self, tmp_path, monkeypatch):
        import base64
        import io
        from PIL import Image

        path = tmp_path / "oriented.jpg"
        exif = Image.Exif()
        exif[274] = 6
        Image.new("RGB", (4, 2), "red").save(path, exif=exif)
        original = path.read_bytes()
        captured = {}

        class Response:
            def raise_for_status(self):
                pass

            def json(self):
                return {"choices": [{"message": {"content": "A red image."}}]}

        monkeypatch.setattr(dl.config, "IMAGE_VISION_BASE_URL", "http://localhost:1234")
        monkeypatch.setattr(dl.config, "IMAGE_VISION_MODEL", "test-model")
        monkeypatch.setattr(dl.requests, "post", lambda _url, **kwargs: captured.update(kwargs) or Response())

        assert dl._summarize_image_with_llm(path) == "A red image."

        payload = captured["json"]
        prompt = payload["messages"][0]["content"][0]["text"]
        assert "Searchable keywords:" in prompt
        encoded = payload["messages"][0]["content"][1]["image_url"]["url"].split(",", 1)[1]
        corrected_bytes = base64.b64decode(encoded)
        with Image.open(io.BytesIO(corrected_bytes)) as corrected:
            assert corrected.size == (2, 4)
            assert corrected.getexif().get(274) in (None, 1)
        assert path.read_bytes() == original

    def test_extract_exif_reads_sub_ifds_and_normalizes_values(self, tmp_path, monkeypatch):
        from PIL import ExifTags, Image

        path = tmp_path / "photo.jpg"
        exif = Image.Exif()
        exif[271] = "Canon"
        exif[274] = 6
        exif.get_ifd(ExifTags.IFD.Exif).update(
            {33434: 0.004, 33437: 2.8, 37385: 25, 34850: 3, 34855: 400, 36867: "2020:06:01 12:00:00"}
        )
        exif.get_ifd(ExifTags.IFD.GPSInfo).update(
            {1: "N", 2: (47.0, 30.0, 0.0), 3: "E", 4: (8.0, 15.0, 36.0)}
        )
        Image.new("RGB", (4, 2)).save(path, exif=exif)
        monkeypatch.setattr(
            dl.config,
            "IMAGE_EXIF_FIELDS",
            ["Make", "Orientation", "ExposureTime", "FNumber", "Flash", "ExposureProgram",
             "ISOSpeedRatings", "DateTimeOriginal", "GPS:GPSLatitude", "GPS:GPSLatitudeRef",
             "GPS:GPSLongitude", "Software"],
        )

        result = dl._extract_image_exif(path)

        assert result == {
            "Make": "Canon",
            "Orientation": 6,
            "ExposureTime": "1/250",
            "FNumber": 2.8,
            "Flash": "flash fired",
            "ExposureProgram": "aperture priority",
            "ISOSpeedRatings": 400,
            "DateTimeOriginal": "2020:06:01 12:00:00",
            "GPS:GPSLatitude": 47.5,
            "GPS:GPSLatitudeRef": "N",
            "GPS:GPSLongitude": 8.26,
        }

    def test_extract_exif_ignores_rationals_with_zero_denominator(self, tmp_path, monkeypatch):
        from PIL import ExifTags, Image
        from PIL.TiffImagePlugin import IFDRational

        path = tmp_path / "zero.jpg"
        exif = Image.Exif()
        exif[271] = "Canon"
        exif.get_ifd(ExifTags.IFD.Exif).update({37382: IFDRational(1, 0), 33437: IFDRational(28, 10)})
        exif.get_ifd(ExifTags.IFD.GPSInfo).update({2: (IFDRational(47, 1), IFDRational(30, 0), IFDRational(0, 0))})
        Image.new("RGB", (4, 2)).save(path, exif=exif)
        monkeypatch.setattr(
            dl.config, "IMAGE_EXIF_FIELDS", ["Make", "SubjectDistance", "FNumber", "GPS:GPSLatitude"]
        )

        assert dl._extract_image_exif(path) == {"Make": "Canon", "FNumber": 2.8}

    @pytest.mark.parametrize("orientation", [None, 1])
    def test_vision_does_not_rotate_image_without_needed_correction(self, tmp_path, monkeypatch, orientation):
        import base64
        from PIL import Image

        path = tmp_path / "already-rotated.jpg"
        image = Image.new("RGB", (2, 4), "blue")
        if orientation is None:
            image.save(path)
        else:
            exif = Image.Exif()
            exif[274] = orientation
            image.save(path, exif=exif)
        original = path.read_bytes()
        captured = {}

        class Response:
            def raise_for_status(self):
                pass

            def json(self):
                return {"choices": [{"message": {"content": "A blue image."}}]}

        monkeypatch.setattr(dl.config, "IMAGE_VISION_BASE_URL", "http://localhost:1234")
        monkeypatch.setattr(dl.config, "IMAGE_VISION_MODEL", "test-model")
        monkeypatch.setattr(dl.requests, "post", lambda _url, **kwargs: captured.update(kwargs) or Response())

        dl._summarize_image_with_llm(path)

        encoded = captured["json"]["messages"][0]["content"][1]["image_url"]["url"].split(",", 1)[1]
        assert base64.b64decode(encoded) == original
        assert path.read_bytes() == original

    def test_unsupported_extension_rejected(self, tmp_path):
        path = tmp_path / "blob.bin"
        path.write_bytes(b"\x00\x01\x02")
        with pytest.raises(ValueError, match="Unsupported"):
            dl.load_document(path)


class TestOCR:
    def test_is_ocr_available_reflects_module_and_binary(self, monkeypatch):
        monkeypatch.setattr(dl, "_has_module", lambda name: name == "pytesseract")
        monkeypatch.setattr(dl, "_has_tesseract_binary", lambda: True)
        assert dl.is_ocr_available() is True

        monkeypatch.setattr(dl, "_has_tesseract_binary", lambda: False)
        assert dl.is_ocr_available() is False

    def test_ocr_keeps_confident_words_and_drops_weak_words(self, monkeypatch):
        from PIL import Image

        class FakePytesseract:
            class Output:
                DICT = "dict"

            @staticmethod
            def image_to_data(_image, *, lang, output_type):
                assert lang == dl.config.OCR_LANGUAGE
                assert output_type == "dict"
                return {
                    "text": ["Service", "gibberish"],
                    "conf": ["94", "18"],
                    "block_num": ["1", "1"],
                    "par_num": ["1", "1"],
                    "line_num": ["1", "1"],
                }

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_configured_tesseract", lambda: FakePytesseract)
        monkeypatch.setattr(dl, "_correct_orientation", lambda _pytesseract, image: image)

        result = dl._ocr_pil_image(Image.new("RGB", (100, 60), color="white"))

        assert result == "Service"

    def test_ocr_rejects_low_confidence_single_word(self, monkeypatch):
        from PIL import Image

        class FakePytesseract:
            class Output:
                DICT = "dict"

            @staticmethod
            def image_to_data(_image, **_kwargs):
                return {"text": ["noise"], "conf": ["62"]}

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_configured_tesseract", lambda: FakePytesseract)
        monkeypatch.setattr(dl, "_correct_orientation", lambda _pytesseract, image: image)

        assert dl._ocr_pil_image(Image.new("RGB", (100, 60), color="white")) == ""

    def test_image_ocr_enabled_uses_ocr_instead_of_vision(self, tmp_path, monkeypatch):
        path = tmp_path / "bill.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n")

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_open_image_rgb", lambda _path: object())
        monkeypatch.setattr(dl, "_ocr_pil_image", lambda _image: "Total: 42.00 EUR")

        def unexpected_call(_path):
            raise AssertionError("OCR image ingestion must not run photo analysis or EXIF extraction")

        monkeypatch.setattr(dl, "_summarize_image_with_llm", unexpected_call)
        monkeypatch.setattr(dl, "_extract_image_exif", unexpected_call)

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "Total: 42.00 EUR" in loaded.content
        assert loaded.metadata["ocr_used"] is True
        assert loaded.metadata["extraction_method"] == "ocr"
        assert "image_exif" not in loaded.metadata

    def test_image_ocr_disabled_keeps_vision_pipeline(self, tmp_path, monkeypatch):
        path = tmp_path / "photo.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n")

        monkeypatch.setattr(dl, "_summarize_image_with_llm", lambda _path, **_kw: "A scenic mountain view.")
        monkeypatch.setattr(dl, "_extract_image_exif", lambda _path: {})

        loaded = dl.load_document(path, ocr_enabled=False)
        assert "A scenic mountain view." in loaded.content
        assert loaded.metadata["ocr_used"] is False

    def test_image_ocr_requested_but_unavailable_falls_back(self, tmp_path, monkeypatch):
        path = tmp_path / "photo.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n")

        monkeypatch.setattr(dl, "is_ocr_available", lambda: False)
        monkeypatch.setattr(dl, "_summarize_image_with_llm", lambda _path, **_kw: "A scenic mountain view.")
        monkeypatch.setattr(dl, "_extract_image_exif", lambda _path: {})

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "A scenic mountain view." in loaded.content
        assert loaded.metadata["ocr_requested_but_unavailable"] is True
        assert loaded.metadata["ocr_used"] is False

    def test_image_ocr_and_vision_analysis_are_mutually_exclusive(self, tmp_path, monkeypatch):
        """OCR and the vision-LLM description are never both run for the same image."""
        path = tmp_path / "photo.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n")

        vision_calls: list[Path] = []
        ocr_calls: list[object] = []
        monkeypatch.setattr(dl, "_summarize_image_with_llm", lambda p, **_kw: vision_calls.append(p) or "A scenic mountain view.")
        monkeypatch.setattr(dl, "_extract_image_exif", lambda _path: {})
        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_open_image_rgb", lambda _path: object())
        monkeypatch.setattr(dl, "_ocr_pil_image", lambda image: ocr_calls.append(image) or "Total: 42.00 EUR")

        # OCR enabled: only the OCR path should run.
        loaded_ocr = dl.load_document(path, ocr_enabled=True)
        assert ocr_calls and not vision_calls
        assert loaded_ocr.metadata["ocr_used"] is True
        assert "image_summary" not in loaded_ocr.metadata

        vision_calls.clear()
        ocr_calls.clear()

        # OCR disabled: only the vision-description path should run.
        loaded_vision = dl.load_document(path, ocr_enabled=False)
        assert vision_calls and not ocr_calls
        assert loaded_vision.metadata["ocr_used"] is False
        assert loaded_vision.metadata["content_kind"] == "image"
        assert "image_summary" in loaded_vision.metadata

    def _make_pdf_with_text(self, path: Path, text: str) -> None:
        import fitz

        document = fitz.open()
        try:
            page = document.new_page()
            page.insert_text((72, 72), text)
            document.save(str(path))
        finally:
            document.close()

    def test_scanned_pdf_triggers_ocr_fallback(self, tmp_path, monkeypatch):
        path = tmp_path / "scan.pdf"
        import fitz

        document = fitz.open()
        try:
            document.new_page()  # blank page: no extractable text layer
            document.save(str(path))
        finally:
            document.close()

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_ocr_image_bytes", lambda _data: "Invoice total: 99.90 EUR")

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "Invoice total: 99.90 EUR" in loaded.content
        assert loaded.metadata["ocr_used"] is True
        assert loaded.metadata["extraction_method"] == "ocr"

    def test_text_pdf_skips_ocr(self, tmp_path, monkeypatch):
        path = tmp_path / "digital.pdf"
        self._make_pdf_with_text(path, "This invoice already has a real embedded text layer.")

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        ocr_calls: list[bytes] = []
        monkeypatch.setattr(dl, "_ocr_image_bytes", lambda data: ocr_calls.append(data) or "SHOULD NOT BE USED")

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "This invoice already has a real embedded text layer." in loaded.content
        assert loaded.metadata["ocr_used"] is False
        assert not ocr_calls

    def test_pdf_text_page_with_background_image_ocrs_only_the_image(self, tmp_path, monkeypatch):
        import fitz
        import io
        from PIL import Image

        path = tmp_path / "mixed.pdf"
        document = fitz.open()
        try:
            page = document.new_page()
            page.insert_text((72, 72), "Real text layer on top of a background photo.")
            buffer = io.BytesIO()
            Image.new("RGB", (120, 120), color="white").save(buffer, format="PNG")
            page.insert_image(fitz.Rect(72, 200, 172, 300), stream=buffer.getvalue())
            document.save(str(path))
        finally:
            document.close()

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_ocr_image_bytes", lambda _data: "Hidden text in background photo")

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "Real text layer on top of a background photo." in loaded.content
        assert "[Image OCR]: Hidden text in background photo" in loaded.content
        assert loaded.metadata["ocr_used"] is True
        assert loaded.metadata["extraction_method"] == "mixed"
        assert loaded.metadata["ocr_image_count"] == 1
        assert loaded.metadata["ocr_pages"] == 0

    def test_pdf_skips_tiny_and_duplicate_embedded_images(self, tmp_path, monkeypatch):
        import fitz
        import io
        from PIL import Image

        path = tmp_path / "duplicate-images.pdf"
        document = fitz.open()
        try:
            page = document.new_page()
            page.insert_text((72, 72), "A usable text layer is present on this page.")
            image_buffers = []
            for size in ((120, 120), (20, 20)):
                buffer = io.BytesIO()
                Image.new("RGB", size, color="white").save(buffer, format="PNG")
                image_buffers.append(buffer.getvalue())
            page.insert_image(fitz.Rect(72, 150, 172, 250), stream=image_buffers[0])
            page.insert_image(fitz.Rect(200, 150, 300, 250), stream=image_buffers[0])
            page.insert_image(fitz.Rect(320, 150, 340, 170), stream=image_buffers[1])
            document.save(str(path))
        finally:
            document.close()

        ocr_calls: list[bytes] = []
        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(
            dl,
            "_ocr_image_bytes",
            lambda data: ocr_calls.append(data) or "Detected label",
        )

        loaded = dl.load_document(path, ocr_enabled=True)

        assert len(ocr_calls) == 1
        assert loaded.metadata["ocr_image_count"] == 1

    def test_ocr_disabled_never_touches_ocr_path(self, tmp_path, monkeypatch):
        path = tmp_path / "scan.pdf"
        import fitz

        document = fitz.open()
        try:
            document.new_page()
            document.save(str(path))
        finally:
            document.close()

        def _boom(_data):
            raise AssertionError("OCR should not run when ocr_enabled=False")

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_ocr_image_bytes", _boom)

        loaded = dl.load_document(path, ocr_enabled=False)
        assert loaded.metadata["ocr_used"] is False

    def test_docx_embedded_image_ocr(self, tmp_path, monkeypatch):
        from docx import Document as DocxDocument
        import io
        from PIL import Image

        path = tmp_path / "letter.docx"
        document = DocxDocument()
        document.add_paragraph("Cover letter text.")
        buffer = io.BytesIO()
        Image.new("RGB", (10, 10), color="white").save(buffer, format="PNG")
        buffer.seek(0)
        document.add_picture(buffer)
        document.save(str(path))

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_ocr_image_bytes", lambda _data: "Scanned signature block")

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "Cover letter text." in loaded.content
        assert "[Image OCR]: Scanned signature block" in loaded.content
        assert loaded.metadata["ocr_used"] is True
        assert loaded.metadata["ocr_image_count"] == 1

    def test_pptx_embedded_image_ocr(self, tmp_path, monkeypatch):
        from pptx import Presentation
        from pptx.util import Inches
        import io
        from PIL import Image

        path = tmp_path / "deck.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        buffer = io.BytesIO()
        Image.new("RGB", (10, 10), color="white").save(buffer, format="PNG")
        buffer.seek(0)
        slide.shapes.add_picture(buffer, Inches(1), Inches(1), Inches(2), Inches(2))
        prs.save(str(path))

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_ocr_image_bytes", lambda _data: "Receipt line items")

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "[Image OCR]: Receipt line items" in loaded.content
        assert loaded.metadata["ocr_used"] is True
        assert loaded.metadata["ocr_image_count"] == 1

    def test_xlsx_embedded_image_ocr(self, tmp_path, monkeypatch):
        import openpyxl
        from openpyxl.drawing.image import Image as XlImage
        import io
        from PIL import Image

        path = tmp_path / "book.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Score"])
        buffer = io.BytesIO()
        Image.new("RGB", (10, 10), color="white").save(buffer, format="PNG")
        buffer.seek(0)
        xl_image = XlImage(buffer)
        ws.add_image(xl_image, "A3")
        wb.save(str(path))

        monkeypatch.setattr(dl, "is_ocr_available", lambda: True)
        monkeypatch.setattr(dl, "_ocr_image_bytes", lambda _data: "Table receipt text")

        loaded = dl.load_document(path, ocr_enabled=True)
        assert "[Image OCR" in loaded.content
        assert "Table receipt text" in loaded.content
        assert loaded.metadata["ocr_used"] is True
        assert loaded.metadata["ocr_image_count"] == 1
